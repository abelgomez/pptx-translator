import io
import logging
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import call, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from pptx_translator.cli import _resolve_log_level, _should_skip_translation
from pptx_translator.config import Settings
from pptx_translator.exceptions_list import (
    ExceptionMode,
    apply_exception_rules,
    load_exception_rules,
    mask_exception_rules,
    restore_exception_rules,
    split_rules_by_mode,
)
from pptx_translator.figure_heuristics import looks_like_figure_label
from pptx_translator.media import _remove_orphaned_timing_nodes
from pptx_translator.pptx_processor import PresentationTranslator
from pptx_translator.translators.base import BaseTranslator, TranslationError
from pptx_translator.translators.factory import create_translator


class ExceptionRulesTests(unittest.TestCase):
    def _rules_from_text(self, text):
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "exceptions.txt"
            path.write_text(text, encoding="utf-8")
            return load_exception_rules(path)

    def test_numeric_wildcards_are_replaced_and_restored(self):
        rules = self._rules_from_text("~Práctica {num} = Practical Lesson {num}\n")
        _, protected = split_rules_by_mode(rules)
        masked, replacements = mask_exception_rules("Práctica 1.", protected)
        self.assertEqual(restore_exception_rules(masked, replacements), "Practical Lesson 1.")
        masked, replacements = mask_exception_rules("práctica 3", protected)
        self.assertEqual(restore_exception_rules(masked, replacements), "Practical Lesson 3")

    def test_default_mode_is_case_insensitive(self):
        rules = self._rules_from_text("SGBD = RDBM\n")
        self.assertEqual(rules[0].mode, ExceptionMode.CASE_INSENSITIVE)

    def test_strict_mode_is_case_sensitive(self):
        rules = self._rules_from_text("!SGA = SGA\n")
        self.assertEqual(rules[0].mode, ExceptionMode.STRICT)
        self.assertIsNone(rules[0].pattern.match("sga"))
        self.assertIsNotNone(rules[0].pattern.match("SGA"))

    def test_pre_translation_mode_uses_plain_substitution(self):
        rules = self._rules_from_text("<Práctica {num} = Practical Lesson {num}\n")
        pre_rules, protected_rules = split_rules_by_mode(rules)
        self.assertEqual(len(pre_rules), 1)
        self.assertEqual(len(protected_rules), 0)
        self.assertEqual(apply_exception_rules("Práctica 1.", pre_rules), "Practical Lesson 1.")

    def test_blank_lines_and_comments_are_ignored(self):
        rules = self._rules_from_text("# comment\n\n   \n<Práctica {num} = Practical Lesson {num}\n")
        self.assertEqual(len(rules), 1)

    def test_placeholder_tokens_are_purely_alphabetic_and_unique(self):
        rules = self._rules_from_text("~Práctica {num} = Practical Lesson {num}\n!SGA = SGA\n")
        masked_a, replacements_a = mask_exception_rules(
            "Práctica 1", rules, start_index=0
        )
        masked_b, replacements_b = mask_exception_rules(
            "SGA y práctica 2", rules, start_index=len(replacements_a)
        )

        self.assertTrue(all(token.isalpha() for token in replacements_a))
        self.assertTrue(all(token.isalpha() for token in replacements_b))
        self.assertNotEqual(set(replacements_a), set(replacements_b))
        self.assertEqual(
            restore_exception_rules(masked_a, replacements_a),
            "Practical Lesson 1",
        )
        self.assertEqual(
            restore_exception_rules(masked_b, replacements_b),
            "SGA y Practical Lesson 2",
        )

    def test_restore_handles_mangled_casing(self):
        rules = self._rules_from_text("!SGA = SGA\n")
        masked, replacements = mask_exception_rules("SGA: System Global Area", rules)
        token = next(iter(replacements))
        capitalized = masked.replace(token, token.capitalize())
        self.assertEqual(
            restore_exception_rules(capitalized, replacements),
            "SGA: System Global Area",
        )

    def test_invalid_line_raises(self):
        with self.assertRaises(ValueError):
            self._rules_from_text("this line has no separator\n")


class CLIInterruptAndLogTests(unittest.TestCase):
    def test_main_reports_abort_without_stacktrace(self):
        import pptx_translator.cli as cli_module

        with patch("pptx_translator.cli.build_arg_parser", side_effect=KeyboardInterrupt):
            with patch("sys.stderr", new_callable=io.StringIO) as stderr:
                exit_code = cli_module.main([])

        self.assertEqual(exit_code, 130)
        self.assertIn("Execution aborted by the user", stderr.getvalue())

    def test_no_flag_uses_configured_log_level(self):
        self.assertEqual(_resolve_log_level(0, "DEBUG"), logging.DEBUG)
        self.assertEqual(_resolve_log_level(0, "INFO"), logging.INFO)
        self.assertEqual(_resolve_log_level(0, "ERROR"), logging.ERROR)
        self.assertEqual(_resolve_log_level(0, "WARNING"), logging.WARNING)

    def test_invalid_configured_level_falls_back_to_warning(self):
        self.assertEqual(_resolve_log_level(0, "NOT_A_LEVEL"), logging.WARNING)

    def test_verbose_flags_override_configured_level(self):
        self.assertEqual(_resolve_log_level(1, "ERROR"), logging.INFO)
        self.assertEqual(_resolve_log_level(2, "ERROR"), logging.DEBUG)

    def test_translation_retries_five_times_with_incremental_backoff(self):
        from pptx_translator.cli import _translate_with_retries

        fake_translator = unittest.mock.Mock()
        fake_translator.translate.side_effect = RuntimeError("temporary failure")

        with patch("pptx_translator.cli.time.sleep") as sleep_mock:
            result = _translate_with_retries(
                unittest.mock.Mock(),
                fake_translator,
                Path("demo.pptx"),
                Path("demo_en.pptx"),
                "en",
                None,
                max_attempts=5,
                initial_delay_seconds=2,
            )

        self.assertIsNone(result)
        self.assertEqual(fake_translator.translate.call_count, 5)
        self.assertEqual(sleep_mock.call_args_list, [call(2), call(4), call(8), call(16)])

    def test_openai_retries_full_slide_batch_before_per_item_fallback(self):
        from pptx_translator.translators.openai import OpenAITranslator

        translator = OpenAITranslator(api_key="demo", api_base_url="https://example.test/v1")
        request_mock = patch.object(
            translator,
            "_request_translation",
            side_effect=[
                "```json\n{bad\n```",
                "```json\n{\"translations\":[{\"id\":\"0\",\"text\":\"TR:Concepto\"},{\"id\":\"1\",\"text\":\"TR:importante\"}]}\n```",
            ],
        )

        with request_mock as request_call, patch.object(
            translator,
            "_translate_single_text",
            side_effect=AssertionError("per-item fallback should not be used"),
        ) as single_call:
            results = translator._translate_slide(["Concepto", " importante"], "es", "en")

        self.assertEqual(results, {"Concepto": "TR:Concepto", " importante": "TR:importante"})
        self.assertEqual(request_call.call_count, 2)
        single_call.assert_not_called()


class PresentationFormattingTests(unittest.TestCase):
    def test_adjacent_runs_with_same_format_are_merged_before_translation(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(self, texts, source_lang, target_lang, context=None, exception_rules=None):
                return {text: "Translated text" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "merge_runs.pptx"
            output_path = Path(tmp_dir) / "merge_runs_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            textbox = slide.shapes.add_textbox(10, 10, 200, 100)
            paragraph = textbox.text_frame.paragraphs[0]
            para_a = paragraph.add_run()
            para_a.text = "Conocer la"
            para_b = paragraph.add_run()
            para_b.text = " estructura"

            prs.save(input_path)
            PresentationTranslator(_StubTranslator()).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            translated = Presentation(output_path)
            out_paragraph = translated.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual(len(out_paragraph.runs), 1)
            self.assertEqual(out_paragraph.text, "Translated text")

    def test_spacing_is_preserved_at_run_boundaries(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(self, texts, source_lang, target_lang, context=None, exception_rules=None):
                return {text: f"TR:{text}" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "boundary_space_test.pptx"
            output_path = Path(tmp_dir) / "boundary_space_test_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            textbox = slide.shapes.add_textbox(10, 10, 300, 100)
            paragraph = textbox.text_frame.paragraphs[0]

            first_run = paragraph.add_run()
            first_run.text = "Conocer "
            first_run.font.bold = True
            second_run = paragraph.add_run()
            second_run.text = "la"
            second_run.font.bold = True

            prs.save(input_path)
            PresentationTranslator(_StubTranslator()).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            output_prs = Presentation(output_path)
            output_paragraph = output_prs.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual(output_paragraph.text, "TR:Conocer la")

    def test_mixed_inline_styles_keep_separate_runs_and_preserve_spacing(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(self, texts, source_lang, target_lang, context=None, exception_rules=None):
                return {text: f"TR:{text}" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "mixed_inline_styles_test.pptx"
            output_path = Path(tmp_dir) / "mixed_inline_styles_test_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            textbox = slide.shapes.add_textbox(10, 10, 300, 100)
            paragraph = textbox.text_frame.paragraphs[0]

            first_run = paragraph.add_run()
            first_run.text = "Concepto"
            first_run.font.bold = True

            second_run = paragraph.add_run()
            second_run.text = " importante"
            second_run.font.italic = True

            third_run = paragraph.add_run()
            third_run.text = " clave"
            third_run.font.bold = True

            prs.save(input_path)
            PresentationTranslator(_StubTranslator()).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            output_prs = Presentation(output_path)
            output_paragraph = output_prs.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual(len(output_paragraph.runs), 3)
            self.assertEqual(output_paragraph.runs[0].text, "TR:Concepto")
            self.assertEqual(output_paragraph.runs[1].text, " TR:importante")
            self.assertEqual(output_paragraph.runs[2].text, " TR:clave")
            self.assertTrue(output_paragraph.runs[0].font.bold)
            self.assertTrue(output_paragraph.runs[1].font.italic)
            self.assertTrue(output_paragraph.runs[2].font.bold)

    def test_flatten_inline_formatting_merges_runs_with_different_styles(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return {text: f"TR:{text}" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "flatten_inline_formatting.pptx"
            output_path = Path(tmp_dir) / "flatten_inline_formatting_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            textbox = slide.shapes.add_textbox(10, 10, 300, 100)
            paragraph = textbox.text_frame.paragraphs[0]

            first_run = paragraph.add_run()
            first_run.text = "Concepto"
            first_run.font.bold = True

            second_run = paragraph.add_run()
            second_run.text = " importante"
            second_run.font.italic = True

            third_run = paragraph.add_run()
            third_run.text = " clave"
            third_run.font.bold = True

            prs.save(input_path)
            PresentationTranslator(_StubTranslator(), flatten_inline_formatting=True).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            output_prs = Presentation(output_path)
            paragraph = output_prs.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual(len(paragraph.runs), 1)
            self.assertEqual(paragraph.text, "TR:Concepto importante clave")

    def test_notes_runs_keep_their_own_formatting(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return {text: f"TR:{text}" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "notes_test.pptx"
            output_path = Path(tmp_dir) / "notes_test_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            notes = slide.notes_slide.notes_text_frame
            paragraph = notes.paragraphs[0]

            first_run = paragraph.add_run()
            first_run.text = "Concepto"
            first_run.font.bold = True

            second_run = paragraph.add_run()
            second_run.text = " importante"
            second_run.font.italic = True

            prs.save(input_path)
            PresentationTranslator(_StubTranslator()).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            output_prs = Presentation(output_path)
            paragraph = output_prs.slides[0].notes_slide.notes_text_frame.paragraphs[0]
            self.assertEqual(paragraph.runs[0].text, "TR:Concepto")
            self.assertTrue(paragraph.runs[0].font.bold)
            self.assertEqual(paragraph.runs[1].text, " TR:importante")
            self.assertTrue(paragraph.runs[1].font.italic)

    def test_regular_text_boxes_enable_word_wrap_and_autosize(self):
        class _StubTranslator(BaseTranslator):
            name = "stub"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return {text: f"TR:{text}" for text in texts}

        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "textbox_test.pptx"
            output_path = Path(tmp_dir) / "textbox_test_en.pptx"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_textbox(10, 10, 200, 100)
            shape.text_frame.text = "Concepto importante para la prueba"

            prs.save(input_path)
            PresentationTranslator(_StubTranslator()).translate(
                str(input_path), str(output_path), target_lang="en", source_lang="es"
            )

            output_prs = Presentation(output_path)
            output_shape = output_prs.slides[0].shapes[0]
            self.assertTrue(output_shape.text_frame.word_wrap)
            self.assertEqual(output_shape.text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)


class FigureHeuristicTests(unittest.TestCase):
    def test_short_multiline_small_box_is_a_figure_caption(self):
        result = looks_like_figure_label(
            "Ordenador con\nsoftware Oracle",
            shape_width_emu=900000,
            shape_height_emu=500000,
            slide_width_emu=12192000,
            slide_height_emu=6858000,
        )
        self.assertTrue(result.is_figure)

    def test_long_sentence_is_not_treated_as_figure(self):
        result = looks_like_figure_label(
            "Instancia Oracle: SGA (System Global Area) + Procesos de fondo. Es el medio para acceder a los datos de la base de datos.",
            shape_width_emu=9000000,
            shape_height_emu=3000000,
            slide_width_emu=12192000,
            slide_height_emu=6858000,
        )
        self.assertFalse(result.is_figure)

    def test_single_line_titles_do_not_become_figures(self):
        result = looks_like_figure_label(
            "Objetivos:",
            shape_width_emu=900000,
            shape_height_emu=300000,
            slide_width_emu=12192000,
            slide_height_emu=6858000,
        )
        self.assertFalse(result.is_figure)


class TranslatorFallbackTests(unittest.TestCase):
    def test_remote_translator_uses_single_item_fallback_after_batch_failure(self):
        class _RemoteLikeTranslator(BaseTranslator):
            name = "openai"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                raise TranslationError("batch failed")

            def _translate_single_text(
                self,
                text,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return f"TR:{text}"

        class _LocalFallback(BaseTranslator):
            name = "local"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return {text: text for text in texts}

            def _translate_single_text(
                self,
                text,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return text

        translator = _RemoteLikeTranslator()
        results, failed = translator.translate_slide(["Práctica"], "es", "en", fallback=_LocalFallback())

        self.assertEqual(results, {"Práctica": "TR:Práctica"})
        self.assertEqual(failed, [])

    def test_local_translator_keeps_original_text_on_single_item_failure(self):
        class _LocalLikeTranslator(BaseTranslator):
            name = "local"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                raise TranslationError("batch failed")

            def _translate_single_text(
                self,
                text,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                raise TranslationError("local item failed")

        class _FallbackTranslator(BaseTranslator):
            name = "fallback"

            def _translate_slide(
                self,
                texts,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return {text: f"TR:{text}" for text in texts}

            def _translate_single_text(
                self,
                text,
                source_lang,
                target_lang,
                context=None,
                exception_rules=None,
            ):
                return f"TR:{text}"

        translator = _LocalLikeTranslator()
        results, failed = translator.translate_slide(["Práctica"], "es", "en", fallback=_FallbackTranslator())

        self.assertEqual(results, {"Práctica": "Práctica"})
        self.assertEqual(failed, ["Práctica"])


class CliBatchProcessingTests(unittest.TestCase):
    def test_should_skip_translation_when_file_already_matches_target_suffix(self):
        self.assertTrue(_should_skip_translation(Path("demo_en.pptx"), "en"))
        self.assertTrue(_should_skip_translation(Path("demo_EN.pptx"), "en"))
        self.assertFalse(_should_skip_translation(Path("demo.pptx"), "en"))
        self.assertFalse(_should_skip_translation(Path("demo_es.pptx"), "en"))

    def test_dry_run_does_not_write_output_or_call_translator(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source = root / "presentation.pptx"
            target = root / "presentation_en.pptx"
            source.write_bytes(b"pptx")
            target.write_bytes(b"existing")

            with patch("pptx_translator.cli.create_translator") as mock_factory, patch(
                "pptx_translator.cli.PresentationTranslator"
            ) as mock_translator_cls:
                mock_factory.return_value.name = "stub"
                result = __import__("pptx_translator.cli", fromlist=["main"]).main(
                    [str(source), "-t", "en", "--dry-run"]
                )

            self.assertEqual(result, 0)
            self.assertEqual(mock_translator_cls.call_count, 0)
            self.assertEqual(target.read_bytes(), b"existing")

    def test_directory_with_recursive_processing_handles_nested_ppts(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            nested_dir = root / "nested"
            nested_dir.mkdir()
            (root / "top.pptx").write_bytes(b"pptx")
            (nested_dir / "inner.pptx").write_bytes(b"pptx")

            with patch("pptx_translator.cli.create_translator") as mock_factory, patch(
                "pptx_translator.cli.PresentationTranslator"
            ) as mock_translator_cls:
                mock_factory.return_value.name = "stub"
                mock_translator_cls.return_value.translate.return_value = type(
                    "Stats",
                    (),
                    {
                        "detected_source_lang": "es",
                        "slides": 1,
                        "audio_removed": 0,
                        "figures_detected": 0,
                        "unique_texts": 1,
                        "translated_units": 1,
                        "failed_texts": 0,
                        "reasons_log": [],
                    },
                )()

                result = __import__("pptx_translator.cli", fromlist=["main"]).main(
                    [str(root), "-t", "en", "-r"]
                )

            self.assertEqual(result, 0)
            self.assertEqual(mock_translator_cls.call_count, 2)


class AudioTimingCleanupTests(unittest.TestCase):
    def _build_slide_with_timing(self, referenced_spid: str, other_spid: str | None = None):
        p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        sld = etree.SubElement(etree.Element(f"{{{p_ns}}}root"), f"{{{p_ns}}}sld")
        timing = etree.SubElement(sld, f"{{{p_ns}}}timing")
        tn_lst = etree.SubElement(timing, f"{{{p_ns}}}tnLst")

        def _add_par_for(spid):
            par = etree.SubElement(tn_lst, f"{{{p_ns}}}par")
            c_tn = etree.SubElement(par, f"{{{p_ns}}}cTn")
            child_tn_lst = etree.SubElement(c_tn, f"{{{p_ns}}}childTnLst")
            cond = etree.SubElement(child_tn_lst, f"{{{p_ns}}}cond")
            tgt_el = etree.SubElement(cond, f"{{{p_ns}}}tgtEl")
            etree.SubElement(tgt_el, f"{{{p_ns}}}spTgt", spid=spid)

        _add_par_for(referenced_spid)
        if other_spid is not None:
            _add_par_for(other_spid)

        return type("FakeSlide", (), {"_element": sld})(), timing

    def test_dangling_timing_entry_is_removed(self):
        slide, _ = self._build_slide_with_timing("5")
        _remove_orphaned_timing_nodes(slide, ["5"])
        self.assertIsNone(slide._element.find(qn("p:timing")))

    def test_unrelated_timing_entries_are_kept(self):
        slide, timing = self._build_slide_with_timing("5", other_spid="7")
        _remove_orphaned_timing_nodes(slide, ["5"])
        remaining = timing.findall(f".//{qn('p:spTgt')}")
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0].attrib.get("spid"), "7")


class TranslatorFactoryTests(unittest.TestCase):
    def test_factory_uses_local_provider_without_configuration(self):
        translator = create_translator(Settings(), provider="local")
        self.assertEqual(translator.name, "local")

    def test_factory_supports_openai_provider(self):
        provider = create_translator(Settings(provider="openai", api_key="demo"), provider="openai")
        self.assertEqual(provider.name, "openai")


if __name__ == "__main__":
    unittest.main()
